import google.generativeai as genai
import io
import streamlit as st
import pandas as pd
import sqlite3
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles.numbers import is_date_format
import xlrd
import sqlparse  # NEW: for SQL formatting

st.title("Excel → Table → Database → SQL Query → Join & Union")

EXCEL_FOLDER = Path("excel_data")
EXCEL_FOLDER.mkdir(exist_ok=True)

# ---------- SIDEBAR: File Uploader ----------
st.sidebar.header("📂 Upload Data")
uploaded_files = st.sidebar.file_uploader(
    "Upload Excel or CSV files",
    type=["xlsx", "xls", "csv"],
    accept_multiple_files=True
)

files = []

if uploaded_files:
    for uploaded_file in uploaded_files:
        file_path = EXCEL_FOLDER / uploaded_file.name
        with open(file_path, "wb") as f:
            f.write(uploaded_file.getbuffer())
        files.append(file_path)
else:
    # Fallback: read from excel_data folder
    files = [
        file
        for file in (
            list(EXCEL_FOLDER.glob("*.xlsx"))
            + list(EXCEL_FOLDER.glob("*.xls"))
            + list(EXCEL_FOLDER.glob("*.csv"))
        )
        if not file.name.startswith("~$")
    ]

# ---------- SIDEBAR: Database Mode ----------
st.sidebar.header("⚙️ Database Settings")
db_mode = st.sidebar.radio(
    "Database Mode",
    ["In-Memory (resets on refresh)", "File-Based (saves data)"]
)

if db_mode == "File-Based (saves data)":
    conn = sqlite3.connect("excel_database.db")
else:
    conn = sqlite3.connect(":memory:")

if "query_history" not in st.session_state:
    st.session_state.query_history = []

tables = {}
excel_types = {}

for file in files:

    # CSV
    if file.suffix.lower() == ".csv":

        df = pd.read_csv(file)

        table_name = file.stem.replace(" ", "_")

        df.to_sql(
            table_name,
            conn,
            if_exists="replace",
            index=False
        )

        tables[table_name] = df

        # CSV has no Excel formatting
        excel_types[table_name] = {
            column: "TEXT"
            for column in df.columns
        }


    # XLS
    elif file.suffix.lower() == ".xls":

        sheets = pd.read_excel(
            file,
            sheet_name=None,
            engine="xlrd"
        )

        workbook = xlrd.open_workbook(
            file,
            formatting_info=True
        )

        for sheet_name, df in sheets.items():

            table_name = (
                f"{file.stem}_{sheet_name}"
                .replace(" ", "_")
            )

            df.to_sql(
                table_name,
                conn,
                if_exists="replace",
                index=False
            )

            tables[table_name] = df

            sheet = workbook.sheet_by_name(
                sheet_name
            )

            type_map = {}

            for col_index, column_name in enumerate(
                df.columns
            ):

                sql_type = "TEXT"

                for row_index in range(
                    1,
                    sheet.nrows
                ):

                    cell = sheet.cell(
                        row_index,
                        col_index
                    )

                    if cell.value == "":
                        continue

                    # DATE
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        sql_type = "DATE"

                    # NUMBER
                    elif cell.ctype == xlrd.XL_CELL_NUMBER:

                        value = cell.value

                        if float(value).is_integer():
                            sql_type = "INTEGER"
                        else:
                            sql_type = "REAL"

                    # TEXT
                    elif cell.ctype == xlrd.XL_CELL_TEXT:
                        sql_type = "TEXT"

                    # UNKNOWN
                    else:
                        sql_type = "TEXT"

                    break

                type_map[column_name] = sql_type

            excel_types[table_name] = type_map


    # XLSX
    elif file.suffix.lower() == ".xlsx":

        sheets = pd.read_excel(
            file,
            sheet_name=None
        )

        workbook = load_workbook(
            file,
            data_only=True
        )

        for sheet_name, df in sheets.items():

            table_name = (
                f"{file.stem}_{sheet_name}"
                .replace(" ", "_")
            )

            df.to_sql(
                table_name,
                conn,
                if_exists="replace",
                index=False
            )

            tables[table_name] = df

            ws = workbook[sheet_name]

            type_map = {}

            for col_index, column_name in enumerate(
                df.columns,
                start=1
            ):

                sql_type = "TEXT"

                for row_index in range(
                    2,
                    ws.max_row + 1
                ):

                    cell = ws.cell(
                        row=row_index,
                        column=col_index
                    )

                    if cell.value is None:
                        continue

                    # DATE
                    if (
                        cell.is_date
                        or is_date_format(
                            cell.number_format
                        )
                    ):
                        sql_type = "DATE"

                    # NUMBER
                    elif cell.data_type == "n":

                        if isinstance(
                            cell.value,
                            int
                        ):
                            sql_type = "INTEGER"
                        else:
                            sql_type = "REAL"

                    # TEXT
                    elif cell.data_type == "s":
                        sql_type = "TEXT"

                    # UNKNOWN
                    else:
                        sql_type = "TEXT"

                    break

                type_map[column_name] = sql_type

            excel_types[table_name] = type_map


# ---------- SQL Formatter Helper ----------
def format_sql(query):
    try:
        return sqlparse.format(query, reindent=True, keyword_case='upper')
    except:
        return query


if not tables:
    st.info("Please add an Excel or CSV file to the 'excel_data' folder.")
    st.stop()


tab1, tab2, tab3, tab4, tab5, tab6, tab7, tab8, tab9, tab10, tab11, tab12, tab13, tab14 = st.tabs(
    [
        "Table",
        "Database",
        "Query",
        "Join & Union",
        "Filter & Sort",
        "Group By",
        "CTE & CASE",
        "Advanced WHERE",
        "Subquery & Window",
        "SQL Functions",
        "SQL Editor",
        "INSERT UPDATE DELETE",
        "Constraints",
        "🤖 AI SQL Generator"
    ]
)

# TABLE
with tab1:

    st.subheader("Excel / CSV Table")

    table_name = st.selectbox(
        "Select Table",
        list(tables.keys()),
        key="table_tab"
    )

    st.dataframe(
        tables[table_name],
        use_container_width=True
    )


# DATABASE
with tab2:

    st.subheader("Database")

    db_tables = pd.read_sql(
        """
        SELECT name
        FROM sqlite_master
        WHERE type='table'
        ORDER BY name
        """,
        conn
    )

    st.write("Database Tables")
    st.dataframe(db_tables, use_container_width=True)

    db_table = st.selectbox(
        "Select Database Table",
        list(tables.keys()),
        key="database_tab"
    )

    database_data = pd.read_sql(
        f"SELECT * FROM {db_table}",
        conn
    )

    st.write("Database Data")
    st.dataframe(
        database_data,
        use_container_width=True
    )

# QUERY
with tab3:

    st.subheader("SQL Query")

    query_table = st.selectbox(
        "Select Table",
        list(tables.keys()),
        key="query_tab"
    )

    df = tables[query_table]

    # 1. CREATE TABLE QUERY
    st.write("### 1. Table / Column Structure Query")

    column_definitions = []

    for column in df.columns:

        sql_type = excel_types[query_table].get(
            column,
            "TEXT"
        )

        column_definitions.append(
            f"`{column}` {sql_type}"
        )

    create_query = f"""
CREATE TABLE `{query_table}` (
    {', '.join(column_definitions)}
);
"""

    st.code(
        create_query,
        language="sql"
    )

    # 2. COLUMN QUERY
    st.write("### 2. Column Query")

    columns_sql = ", ".join(
        f"`{column}`"
        for column in df.columns
    )

    column_query = (
        f"SELECT {columns_sql} "
        f"FROM `{query_table}`;"
    )

    st.code(
        column_query,
        language="sql"
    )

    # 3. DATA QUERY
    st.write("### 3. Data Query")

    def sql_value(value):

        if pd.isna(value):
            return "NULL"

        if isinstance(value, pd.Timestamp):
            return f"'{value.strftime('%Y-%m-%d')}'"

        if isinstance(value, str):
            value = value.replace("'", "''")
            return f"'{value}'"

        return str(value)

    columns = ", ".join(
        f"`{column}`"
        for column in df.columns
    )

    insert_queries = []

    for _, row in df.iterrows():

        values = [
            sql_value(value)
            for value in row
        ]

        insert_query = (
            f"INSERT INTO `{query_table}` "
            f"({columns}) VALUES "
            f"({', '.join(values)});"
        )

        insert_queries.append(insert_query)

    data_query = "\n".join(insert_queries)

    st.code(
        data_query,
        language="sql"
    )

    # 4. RUN QUERY
    st.write("### 4. Run Query")

    if st.button("Run Query"):

        result = pd.read_sql(
            f"SELECT * FROM `{query_table}`",
            conn
        )

        st.dataframe(
            result,
            use_container_width=True
        )

# JOIN & UNION
with tab4:

    st.subheader("Join & Union")

    operations = [
        "INNER JOIN",
        "LEFT JOIN",
        "RIGHT JOIN",
        "FULL OUTER JOIN",
        "CROSS JOIN",
        "SELF JOIN",
        "UNION",
        "UNION ALL",
        "INTERSECT",
        "EXCEPT"
    ]

    operation = st.selectbox(
        "Select Operation",
        operations
    )

    table_list = list(tables.keys())

    left_table = st.selectbox(
        "Select First Table",
        table_list,
        key="join_left_table"
    )

    is_set_operation = operation in [
        "UNION",
        "UNION ALL",
        "INTERSECT",
        "EXCEPT"
    ]

    
    is_set_operation = operation in [
        "UNION",
        "UNION ALL",
        "INTERSECT",
        "EXCEPT"
    ]

    # Default
    can_normal_join = True

    # SELF JOIN
    if operation == "SELF JOIN":

        right_table = left_table

    # UNION / INTERSECT / EXCEPT
    elif is_set_operation:

        if len(table_list) < 2:
            st.warning(
                "Set operation ke liye kam se kam 2 tables chahiye."
            )
            can_normal_join = False
            right_table = left_table

        else:
            right_table = st.selectbox(
                "Select Second Table",
                table_list,
                key="set_right_table"
            )

    # Normal JOIN
    else:

        other_tables = [
            table
            for table in table_list
            if table != left_table
        ]

        if not other_tables:

            st.warning(
                "Join ke liye 2 different tables chahiye. "
                "Dusri Excel/CSV file upload karo."
            )

            can_normal_join = False
            right_table = left_table

        else:

            right_table = st.selectbox(
                "Select Second Table",
                other_tables,
                key="join_right_table"
            )

    left_df = tables[left_table]
    right_df = tables[right_table]


    def quote_column(column):
        return f'"{str(column).replace(chr(34), chr(34) * 2)}"'


    # ---------------- JOIN ----------------

    if not is_set_operation and operation != "CROSS JOIN":

        common_columns = [
            column
            for column in left_df.columns
            if column in right_df.columns
        ]

        if common_columns:

            st.info(
                "Common column found: "
                + ", ".join(map(str, common_columns))
            )

            default_left = 0
            default_right = 0

            left_index = st.selectbox(
                "First Table Column",
                list(left_df.columns),
                index=default_left,
                key="left_column"
            )

            right_index = st.selectbox(
                "Second Table Column",
                list(right_df.columns),
                index=list(right_df.columns).index(
                    common_columns[0]
                ),
                key="right_column"
            )

        else:

            st.warning(
                "No common column found. "
                "Select matching columns manually."
            )

            left_index = st.selectbox(
                "First Table Column",
                list(left_df.columns),
                key="left_column"
            )

            right_index = st.selectbox(
                "Second Table Column",
                list(right_df.columns),
                key="right_column"
            )


        if can_normal_join and st.button("Generate Join Query"):

            left_select = []
            right_select = []

            if operation == "SELF JOIN":

                for column in left_df.columns:

                    left_select.append(
                        f'a.{quote_column(column)} '
                        f'AS {quote_column(column)}'
                    )

                for column in right_df.columns:

                    right_select.append(
                        f'b.{quote_column(column)} '
                        f'AS {quote_column(str(column) + "_B")}'
                    )

                select_sql = ", ".join(
                    left_select + right_select
                )

                query = f"""
SELECT {select_sql}
FROM {quote_column(left_table)} AS a
JOIN {quote_column(left_table)} AS b
ON a.{quote_column(left_index)}
 = b.{quote_column(right_index)};
"""

            else:

                for column in left_df.columns:

                    left_select.append(
                        f'l.{quote_column(column)} '
                        f'AS {quote_column(column)}'
                    )

                for column in right_df.columns:

                    if column in left_df.columns:

                        output_name = (
                            f"{right_table}_{column}"
                        )

                    else:

                        output_name = column

                    right_select.append(
                        f'r.{quote_column(column)} '
                        f'AS {quote_column(output_name)}'
                    )

                select_sql = ", ".join(
                    left_select + right_select
                )

                                # SQLite older versions don't support RIGHT/FULL OUTER JOIN
                if operation == "RIGHT JOIN":
                    # Simulate RIGHT JOIN by swapping tables and using LEFT JOIN
                    query = f"""
SELECT {select_sql}
FROM {quote_column(right_table)} AS l
LEFT JOIN {quote_column(left_table)} AS r
ON l.{quote_column(right_index)}
 = r.{quote_column(left_index)};
"""
                elif operation == "FULL OUTER JOIN":
                    # Simulate FULL OUTER JOIN using UNION of two LEFT JOINs
                    query = f"""
SELECT {select_sql}
FROM {quote_column(left_table)} AS l
LEFT JOIN {quote_column(right_table)} AS r
ON l.{quote_column(left_index)} = r.{quote_column(right_index)}

UNION

SELECT {select_sql}
FROM {quote_column(right_table)} AS l
LEFT JOIN {quote_column(left_table)} AS r
ON l.{quote_column(right_index)} = r.{quote_column(left_index)};
"""
                else:
                    query = f"""
SELECT {select_sql}
FROM {quote_column(left_table)} AS l
{operation}
{quote_column(right_table)} AS r
ON l.{quote_column(left_index)}
 = r.{quote_column(right_index)};
"""


            st.write("### SQL Query")

            st.code(
                query,
                language="sql"
            )

            try:

                result = pd.read_sql(
                    query,
                    conn
                )

                st.write("### Result")

                st.dataframe(
                    result,
                    use_container_width=True
                )

            except Exception as e:

                st.error(
                    f"Query Error: {e}"
                )


    # ---------------- CROSS JOIN ----------------

    elif operation == "CROSS JOIN":

        if st.button("Generate Cross Join Query"):

            left_select = []

            right_select = []

            for column in left_df.columns:

                left_select.append(
                    f'l.{quote_column(column)} '
                    f'AS {quote_column(column)}'
                )

            for column in right_df.columns:

                output_name = (
                    f"{right_table}_{column}"
                    if column in left_df.columns
                    else column
                )

                right_select.append(
                    f'r.{quote_column(column)} '
                    f'AS {quote_column(output_name)}'
                )

            select_sql = ", ".join(
                left_select + right_select
            )

            query = f"""
SELECT {select_sql}
FROM {quote_column(left_table)} AS l
CROSS JOIN {quote_column(right_table)} AS r;
"""

            st.write("### SQL Query")

            st.code(
                query,
                language="sql"
            )

            try:

                result = pd.read_sql(
                    query,
                    conn
                )

                st.write("### Result")

                st.dataframe(
                    result,
                    use_container_width=True
                )

            except Exception as e:

                st.error(
                    f"Query Error: {e}"
                )


    # ---------------- UNION ----------------

    else:

        left_columns = st.multiselect(
            "Select Columns From First Table",
            list(left_df.columns),
            default=list(left_df.columns),
            key="union_left_columns"
        )

        right_columns = st.multiselect(
            "Select Columns From Second Table",
            list(right_df.columns),
            default=list(right_df.columns),
            key="union_right_columns"
        )

        if st.button("Generate Set Query"):

            if len(left_columns) != len(right_columns):

                st.error(
                    "Both tables must have the same number of columns."
                )

            elif len(left_columns) == 0:

                st.error(
                    "Please select at least one column."
                )

            else:

                left_sql = ", ".join(
                    quote_column(column)
                    for column in left_columns
                )

                right_sql = ", ".join(
                    quote_column(column)
                    for column in right_columns
                )

                query = f"""
SELECT {left_sql}
FROM {quote_column(left_table)}

{operation}

SELECT {right_sql}
FROM {quote_column(right_table)};
"""

                st.write("### SQL Query")

                st.code(
                    query,
                    language="sql"
                )

                try:

                    result = pd.read_sql(
                        query,
                        conn
                    )

                    st.write("### Result")

                    st.dataframe(
                        result,
                        use_container_width=True
                    )

                except Exception as e:

                    st.error(
                        f"Query Error: {e}"
                    )

# FILTER & SORT
with tab5:

    st.subheader("Filter & Sort")

    filter_table = st.selectbox(
        "Select Table",
        list(tables.keys()),
        key="filter_table"
    )

    df = tables[filter_table]

    filter_column = st.selectbox(
        "Select Column",
        list(df.columns),
        key="filter_column"
    )

    operator = st.selectbox(
        "Select Operator",
        ["=", "!=", ">", "<", ">=", "<=", "LIKE"],
        key="filter_operator"
    )

    value = st.text_input(
        "Enter Value",
        key="filter_value"
    )

    sort_column = st.selectbox(
        "Sort By",
        list(df.columns),
        key="sort_column"
    )

    sort_order = st.selectbox(
        "Sort Order",
        ["ASC", "DESC"],
        key="sort_order"
    )

    limit = st.number_input(
        "Limit",
        min_value=1,
        value=10,
        key="filter_limit"
    )

    if st.button("Generate Filter Query"):

        if operator == "LIKE":
            condition = f"`{filter_column}` LIKE '%{value}%'"
        else:
            condition = f"`{filter_column}` {operator} '{value}'"

        query = f"""
SELECT *
FROM `{filter_table}`
WHERE {condition}
ORDER BY `{sort_column}` {sort_order}
LIMIT {limit};
"""

        st.write("### SQL Query")

        st.code(
            query,
            language="sql"
        )

        try:

            result = pd.read_sql(
                query,
                conn
            )

            st.write("### Result")

            st.dataframe(
                result,
                use_container_width=True
            )

        except Exception as e:

            st.error(
                f"Query Error: {e}"
            )

# GROUP BY
with tab6:

    st.subheader("GROUP BY & Aggregate")

    group_table = st.selectbox(
        "Select Table",
        list(tables.keys()),
        key="group_table"
    )

    df = tables[group_table]

    group_column = st.selectbox(
        "Group By Column",
        list(df.columns),
        key="group_column"
    )

    aggregate_function = st.selectbox(
        "Aggregate Function",
        ["COUNT", "SUM", "AVG", "MIN", "MAX"],
        key="aggregate_function"
    )

    aggregate_column = st.selectbox(
        "Aggregate Column",
        list(df.columns),
        key="aggregate_column"
    )

    use_having = st.checkbox(
        "Use HAVING",
        key="use_having"
    )

    having_value = st.number_input(
        "HAVING Value",
        value=0,
        key="having_value"
    )

    if st.button("Generate GROUP BY Query"):

        query = f"""
SELECT `{group_column}`,
       {aggregate_function}(`{aggregate_column}`) AS result
FROM `{group_table}`
GROUP BY `{group_column}`
"""

        if use_having:
            query += f"""
HAVING {aggregate_function}(`{aggregate_column}`) > {having_value}
"""

        query += ";"

        st.write("### SQL Query")

        st.code(
            query,
            language="sql"
        )

        try:

            result = pd.read_sql(
                query,
                conn
            )

            st.write("### Result")

            st.dataframe(
                result,
                use_container_width=True
            )

        except Exception as e:

            st.error(
                f"Query Error: {e}"
            )

# CTE & CASE
with tab7:

    st.subheader("CTE & CASE")

    cte_table = st.selectbox(
        "Select Table",
        list(tables.keys()),
        key="cte_table"
    )

    df = tables[cte_table]

    cte_column = st.selectbox(
        "CTE Filter Column",
        list(df.columns),
        key="cte_column"
    )

    cte_operator = st.selectbox(
        "CTE Operator",
        ["=", "!=", ">", "<", ">=", "<="],
        key="cte_operator"
    )

    cte_value = st.text_input(
        "CTE Filter Value",
        key="cte_value"
    )

    case_column = st.selectbox(
        "CASE Column",
        list(df.columns),
        key="case_column"
    )

    case_high = st.number_input(
        "High Value",
        value=50000,
        key="case_high"
    )

    case_medium = st.number_input(
        "Medium Value",
        value=10000,
        key="case_medium"
    )

    if st.button("Generate CTE & CASE Query"):

        query = f"""
WITH filtered_data AS (
    SELECT *
    FROM `{cte_table}`
    WHERE `{cte_column}` {cte_operator} '{cte_value}'
)
SELECT
    *,
    CASE
        WHEN `{case_column}` >= {case_high} THEN 'High'
        WHEN `{case_column}` >= {case_medium} THEN 'Medium'
        ELSE 'Low'
    END AS Category
FROM filtered_data;
"""

        st.write("### SQL Query")

        st.code(
            query,
            language="sql"
        )

        try:

            result = pd.read_sql(
                query,
                conn
            )

            st.write("### Result")

            st.dataframe(
                result,
                use_container_width=True
            )

        except Exception as e:

            st.error(
                f"Query Error: {e}"
            )

# ADVANCED WHERE
with tab8:

    st.subheader("DISTINCT & Advanced WHERE")

    where_table = st.selectbox(
        "Select Table",
        list(tables.keys()),
        key="where_table"
    )

    df = tables[where_table]

    use_distinct = st.checkbox(
        "Use DISTINCT",
        key="use_distinct"
    )

    selected_columns = st.multiselect(
        "Select Columns",
        list(df.columns),
        default=list(df.columns),
        key="where_columns"
    )

    condition_type = st.selectbox(
        "Condition Type",
        [
            "None",
            "AND",
            "OR",
            "IN",
            "BETWEEN"
        ],
        key="condition_type"
    )

    column1 = st.selectbox(
        "Column",
        list(df.columns),
        key="condition_column1"
    )

    operator = st.selectbox(
        "Operator",
        ["=", "!=", ">", "<", ">=", "<="],
        key="condition_operator"
    )

    value1 = st.text_input(
        "Value",
        key="condition_value1"
    )

    value2 = ""

    if condition_type in ["AND", "OR"]:

        column2 = st.selectbox(
            "Second Column",
            list(df.columns),
            key="condition_column2"
        )

        operator2 = st.selectbox(
            "Second Operator",
            ["=", "!=", ">", "<", ">=", "<="],
            key="condition_operator2"
        )

        value2 = st.text_input(
            "Second Value",
            key="condition_value2"
        )

    elif condition_type == "IN":

        value2 = st.text_input(
            "IN Values (comma separated)",
            key="in_values"
        )

    elif condition_type == "BETWEEN":

        value2 = st.text_input(
            "Second Value",
            key="between_value"
        )

    if st.button("Generate WHERE Query"):

        if not selected_columns:

            st.error("Please select at least one column.")

        else:

            select_word = "SELECT DISTINCT" if use_distinct else "SELECT"

            columns_sql = ", ".join(
                f"`{column}`"
                for column in selected_columns
            )

            query = f"{select_word} {columns_sql}\nFROM `{where_table}`"

            if condition_type == "None":

                query += ";"

            elif condition_type in ["AND", "OR"]:

                query += f"""
WHERE `{column1}` {operator} '{value1}'
{condition_type}
`{column2}` {operator2} '{value2}';
"""

            elif condition_type == "IN":

                values = [
                    f"'{value.strip()}'"
                    for value in value2.split(",")
                ]

                query += f"""
WHERE `{column1}` IN ({', '.join(values)});
"""

            elif condition_type == "BETWEEN":

                query += f"""
WHERE `{column1}` BETWEEN '{value1}' AND '{value2}';
"""


            st.write("### SQL Query")

            st.code(
                query,
                language="sql"
            )

            try:

                result = pd.read_sql(
                    query,
                    conn
                )

                st.write("### Result")

                st.dataframe(
                    result,
                    use_container_width=True
                )

            except Exception as e:

                st.error(
                    f"Query Error: {e}"
                )

# SUBQUERY & WINDOW FUNCTIONS
with tab9:

    st.subheader("Subquery & Window Functions")

    feature = st.selectbox(
        "Select Feature",
        [
            "SUBQUERY",
            "ROW_NUMBER",
            "RANK",
            "DENSE_RANK"
        ],
        key="subquery_window_feature"
    )

    table_list = list(tables.keys())

    selected_table = st.selectbox(
        "Select Table",
        table_list,
        key="subquery_window_table"
    )

    df = tables[selected_table]


    # ---------------- SUBQUERY ----------------

    if feature == "SUBQUERY":

        selected_columns = st.multiselect(
            "Select Columns",
            list(df.columns),
            default=list(df.columns),
            key="subquery_columns"
        )

        aggregate_column = st.selectbox(
            "Subquery Column",
            list(df.columns),
            key="subquery_aggregate_column"
        )

        aggregate_function = st.selectbox(
            "Aggregate Function",
            ["AVG", "SUM", "MIN", "MAX", "COUNT"],
            key="subquery_aggregate_function"
        )

        operator = st.selectbox(
            "Operator",
            [">", "<", "=", ">=", "<="],
            key="subquery_operator"
        )

        value_column = st.selectbox(
            "Main Query Column",
            list(df.columns),
            key="subquery_value_column"
        )

        if st.button("Generate Subquery"):

            if not selected_columns:

                st.error("Please select at least one column.")

            else:

                columns_sql = ", ".join(
                    f"`{column}`"
                    for column in selected_columns
                )

                query = f"""
SELECT {columns_sql}
FROM `{selected_table}`
WHERE `{value_column}` {operator}
(
    SELECT {aggregate_function}(`{aggregate_column}`)
    FROM `{selected_table}`
);
"""

                st.write("### SQL Query")

                st.code(
                    query,
                    language="sql"
                )

                try:

                    result = pd.read_sql(
                        query,
                        conn
                    )

                    st.write("### Result")

                    st.dataframe(
                        result,
                        use_container_width=True
                    )

                except Exception as e:

                    st.error(
                        f"Query Error: {e}"
                    )


    # ---------------- WINDOW FUNCTIONS ----------------

    else:

        order_column = st.selectbox(
            "Order By Column",
            list(df.columns),
            key="window_order_column"
        )

        partition_column = st.selectbox(
            "Partition By Column",
            ["None"] + list(df.columns),
            key="window_partition_column"
        )

        if st.button("Generate Window Query"):

            if partition_column == "None":

                query = f"""
SELECT *,
       {feature}() OVER (
           ORDER BY `{order_column}`
       ) AS Window_Result
FROM `{selected_table}`;
"""

            else:

                query = f"""
SELECT *,
       {feature}() OVER (
           PARTITION BY `{partition_column}`
           ORDER BY `{order_column}`
       ) AS Window_Result
FROM `{selected_table}`;
"""

            st.write("### SQL Query")

            st.code(
                query,
                language="sql"
            )

            try:

                result = pd.read_sql(
                    query,
                    conn
                )

                st.write("### Result")

                st.dataframe(
                    result,
                    use_container_width=True
                )

            except Exception as e:

                st.error(
                    f"Query Error: {e}"
                )


# SQL FUNCTIONS
with tab10:

    st.subheader("SQL Functions")

    function = st.selectbox(
        "Select Function",
        [
            "YEAR",
            "MONTH",
            "DAY",
            "UPPER",
            "LOWER",
            "LENGTH",
            "COALESCE",
            "NULLIF"
        ],
        key="sql_function"
    )

    function_table = st.selectbox(
        "Select Table",
        list(tables.keys()),
        key="function_table"
    )

    df = tables[function_table]

    function_column = st.selectbox(
        "Select Column",
        list(df.columns),
        key="function_column"
    )

    if function == "COALESCE":

        default_value = st.text_input(
            "Default Value",
            key="coalesce_value"
        )

    elif function == "NULLIF":

        compare_value = st.text_input(
            "Compare Value",
            key="nullif_value"
        )

    if st.button("Generate Function Query"):

        column = f"`{function_column}`"

        if function == "YEAR":
            expression = f"CAST(strftime('%Y', {column}) AS INTEGER)"

        elif function == "MONTH":
            expression = f"CAST(strftime('%m', {column}) AS INTEGER)"

        elif function == "DAY":
            expression = f"CAST(strftime('%d', {column}) AS INTEGER)"

        elif function == "UPPER":
            expression = f"UPPER({column})"

        elif function == "LOWER":
            expression = f"LOWER({column})"

        elif function == "LENGTH":
            expression = f"LENGTH({column})"

        elif function == "COALESCE":
            expression = f"COALESCE({column}, '{default_value}')"

        else:
            expression = f"NULLIF({column}, '{compare_value}')"

        query = f"""
SELECT
    {column},
    {expression} AS result
FROM `{function_table}`;
"""

        st.write("### SQL Query")

        st.code(
            query,
            language="sql"
        )

        try:

            result = pd.read_sql(
                query,
                conn
            )

            st.write("### Result")

            st.dataframe(
                result,
                use_container_width=True
            )

        except Exception as e:

            st.error(
                f"Query Error: {e}"
            )



# SQL EDITOR

with tab11:

    st.subheader("SQL Query Editor")

    table_list = list(tables.keys())

    st.write("Available Tables")

    st.dataframe(
        pd.DataFrame({"Table Name": table_list}),
        use_container_width=True
    )

    editor_table = st.selectbox(
        "Select Table",
        table_list,
        key="editor_table"
    )

    sql_query = st.text_area(
        "SQL Query",
        value=f"SELECT * FROM `{editor_table}`;",
        height=200,
        key="editor_query"
    )

    if st.button("Execute SQL Query"):

        if not sql_query.strip():

            st.warning("Please enter a SQL query.")

        else:

            try:

                result = pd.read_sql(
                    sql_query,
                    conn
                )

                # Save query to history
                st.session_state.query_history.append(
                    sql_query
                )

                st.write("### Query Result")

                st.dataframe(
                    result,
                    use_container_width=True
                )

                # CSV Download
                csv_data = result.to_csv(index=False)

                st.download_button(
                    "Download CSV",
                    csv_data,
                    "query_result.csv",
                    "text/csv"
                )

                # Excel Download
                excel_buffer = io.BytesIO()

                with pd.ExcelWriter(
                    excel_buffer,
                    engine="openpyxl"
                ) as writer:

                    result.to_excel(
                        writer,
                        index=False,
                        sheet_name="Result"
                    )

                st.download_button(
                    "Download Excel",
                    excel_buffer.getvalue(),
                    "query_result.xlsx",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )

            except Exception as e:

                st.error(
                    f"SQL Error: {e}"
                )


    st.write("### Query History")

    if st.session_state.query_history:

        for i, old_query in enumerate(
            reversed(st.session_state.query_history),
            start=1
        ):

            st.code(
                old_query,
                language="sql"
            )

    else:

        st.info("No query history yet.")


# INSERT UPDATE DELETE

# INSERT UPDATE DELETE

with tab12:

    st.subheader("INSERT / UPDATE / DELETE")

    table_list = list(tables.keys())

    dml_table = st.selectbox(
        "Select Table",
        table_list,
        key="dml_table"
    )

    df = tables[dml_table]

    operation = st.selectbox(
        "Select Operation",
        [
            "INSERT",
            "UPDATE",
            "DELETE"
        ],
        key="dml_operation"
    )

    def sql_value(value, column):

        if value is None:
            return "NULL"

        value = str(value).strip()

        if value == "":
            return "NULL"

        sql_type = excel_types[dml_table].get(
            column,
            "TEXT"
        )

        if sql_type == "INTEGER":

            try:
                return str(int(float(value)))
            except:
                pass

        elif sql_type == "REAL":

            try:
                return str(float(value))
            except:
                pass

        value = value.replace("'", "''")

        return f"'{value}'"

    # INSERT
    if operation == "INSERT":

        insert_columns = st.multiselect(
            "Select Columns",
            list(df.columns),
            default=list(df.columns),
            key="insert_columns"
        )

        insert_values = {}

        for column in insert_columns:

            insert_values[column] = st.text_input(
                f"Value for {column}",
                key=f"insert_value_{column}"
            )

        # STEP 1: Generate query and store it
        if st.button("Generate INSERT Query"):

            if not insert_columns:

                st.error("Please select at least one column.")

            else:

                columns_sql = ", ".join(
                    f"`{column}`"
                    for column in insert_columns
                )

                values_sql = ", ".join(
                    sql_value(insert_values[column], column)
                    for column in insert_columns
                )

                query = f"INSERT INTO `{dml_table}` ({columns_sql}) VALUES ({values_sql});"

                st.session_state["pending_dml_query"] = query

                st.write("### SQL Query")
                st.code(query, language="sql")

        # STEP 2: Execute the stored query
        if "pending_dml_query" in st.session_state:

            if st.button("Execute INSERT"):

                query = st.session_state["pending_dml_query"]

                try:

                    cursor = conn.execute(query)
                    conn.commit()

                    tables[dml_table] = pd.read_sql(
                        f"SELECT * FROM `{dml_table}`",
                        conn
                    )

                    st.session_state.query_history.append(query)

                    st.success(
                        f"INSERT successful. Rows affected: {cursor.rowcount}"
                    )

                    st.write("### Updated Data")
                    st.dataframe(tables[dml_table], use_container_width=True)

                    del st.session_state["pending_dml_query"]

                except Exception as e:

                    st.error(f"SQL Error: {e}")

    # UPDATE
    elif operation == "UPDATE":

        update_column = st.selectbox(
            "Column to Update",
            list(df.columns),
            key="update_column"
        )

        new_value = st.text_input(
            "New Value",
            key="update_value"
        )

        where_column = st.selectbox(
            "WHERE Column",
            list(df.columns),
            key="update_where_column"
        )

        where_operator = st.selectbox(
            "WHERE Operator",
            ["=", "!=", ">", "<", ">=", "<=", "LIKE"],
            key="update_where_operator"
        )

        where_value = st.text_input(
            "WHERE Value",
            key="update_where_value"
        )

        # STEP 1: Generate query and store it
        if st.button("Generate UPDATE Query"):

            if where_value.strip() == "":

                st.error("WHERE value is required.")

            else:

                update_sql = sql_value(new_value, update_column)

                where_sql = sql_value(where_value, where_column)

                if where_operator == "LIKE":

                    where_sql = f"'%{where_value}%'"

                query = f"""
UPDATE `{dml_table}`
SET `{update_column}` = {update_sql}
WHERE `{where_column}` {where_operator} {where_sql};
"""

                st.session_state["pending_dml_query"] = query

                st.write("### SQL Query")
                st.code(query, language="sql")

        # STEP 2: Execute button OUTSIDE the generate block
        if "pending_dml_query" in st.session_state:

            if st.button("Execute UPDATE"):

                query = st.session_state["pending_dml_query"]

                try:

                    cursor = conn.execute(query)
                    conn.commit()

                    tables[dml_table] = pd.read_sql(
                        f"SELECT * FROM `{dml_table}`",
                        conn
                    )

                    st.session_state.query_history.append(query)

                    st.success(
                        f"UPDATE successful. Rows affected: {cursor.rowcount}"
                    )

                    st.write("### Updated Data")
                    st.dataframe(tables[dml_table], use_container_width=True)

                    del st.session_state["pending_dml_query"]

                except Exception as e:

                    st.error(f"SQL Error: {e}")

    # DELETE
    else:

        where_column = st.selectbox(
            "WHERE Column",
            list(df.columns),
            key="delete_where_column"
        )

        where_operator = st.selectbox(
            "WHERE Operator",
            ["=", "!=", ">", "<", ">=", "<=", "LIKE"],
            key="delete_where_operator"
        )

        where_value = st.text_input(
            "WHERE Value",
            key="delete_where_value"
        )

        # STEP 1: Generate query and store it
        if st.button("Generate DELETE Query"):

            if where_value.strip() == "":

                st.error("WHERE value is required.")

            else:

                where_sql = sql_value(where_value, where_column)

                if where_operator == "LIKE":

                    where_sql = f"'%{where_value}%'"

                query = f"""
DELETE FROM `{dml_table}`
WHERE `{where_column}` {where_operator} {where_sql};
"""

                st.session_state["pending_dml_query"] = query

                st.write("### SQL Query")
                st.code(query, language="sql")

        # STEP 2: Execute button OUTSIDE the generate block
        if "pending_dml_query" in st.session_state:

            if st.button("Execute DELETE"):

                query = st.session_state["pending_dml_query"]

                try:

                    cursor = conn.execute(query)
                    conn.commit()

                    tables[dml_table] = pd.read_sql(
                        f"SELECT * FROM `{dml_table}`",
                        conn
                    )

                    st.session_state.query_history.append(query)

                    st.success(
                        f"DELETE successful. Rows affected: {cursor.rowcount}"
                    )

                    st.write("### Updated Data")
                    st.dataframe(tables[dml_table], use_container_width=True)

                    del st.session_state["pending_dml_query"]

                except Exception as e:

                    st.error(f"SQL Error: {e}")

# CONSTRAINTS

with tab13:

    st.subheader("SQL Constraints")

    constraint_table = st.selectbox(
        "Select Table",
        list(tables.keys()),
        key="constraint_table"
    )

    df = tables[constraint_table]

    constraint_column = st.selectbox(
        "Select Column",
        list(df.columns),
        key="constraint_column"
    )

    constraint_type = st.selectbox(
        "Select Constraint",
        [
            "PRIMARY KEY",
            "NOT NULL",
            "UNIQUE",
            "DEFAULT"
        ],
        key="constraint_type"
    )

    default_value = ""

    if constraint_type == "DEFAULT":

        default_value = st.text_input(
            "Default Value",
            key="constraint_default"
        )

    if st.button("Generate Constraint Query"):

        column_types = []

        for column in df.columns:

            sql_type = excel_types[constraint_table].get(
                column,
                "TEXT"
            )

            constraint = ""

            if column == constraint_column:

                if constraint_type == "PRIMARY KEY":
                    constraint = " PRIMARY KEY"

                elif constraint_type == "NOT NULL":
                    constraint = " NOT NULL"

                elif constraint_type == "UNIQUE":
                    constraint = " UNIQUE"

                elif constraint_type == "DEFAULT":
                    constraint = f" DEFAULT '{default_value}'"

            column_types.append(
                f"`{column}` {sql_type}{constraint}"
            )

        query = f"""
CREATE TABLE `new_{constraint_table}` (
    {', '.join(column_types)}
);
"""

        st.write("### SQL Query")

        st.code(
            query,
            language="sql"
        )

# AI SQL GENERATOR

with tab14:

    st.write("✅ AI Tab loaded successfully.")  # Debug line

    st.subheader("🤖 AI SQL Generator (Text to SQL)")

    if "saved_api_key" not in st.session_state:
        st.session_state.saved_api_key = ""

    api_key = st.text_input("Enter Google Gemini API Key", type="password", value=st.session_state.saved_api_key)

    if api_key:
        st.session_state.saved_api_key = api_key

    user_prompt = st.text_area(
        "What do you want to query?",
        placeholder="e.g., Show me the top 5 customers by total sales",
        key="ai_prompt"
    )

    if st.button("Generate SQL with AI"):

        if not api_key:
            st.error("Please enter your Gemini API Key first.")
        elif not user_prompt.strip():
            st.error("Please enter your question.")
        else:
            genai.configure(api_key=api_key)

            schema_info = ""
            for tbl_name, tbl_df in tables.items():
                schema_info += f"Table: {tbl_name}\n"
                schema_info += f"Columns: {', '.join(map(str, tbl_df.columns))}\n\n"

            full_prompt = f"""
You are an expert SQLite developer.
Generate a valid SQLite query for the user's request based on the schema below.
Return ONLY the SQL query, no explanation, no markdown.

Schema:
{schema_info}

User Request: {user_prompt}
"""

            try:
                model = genai.GenerativeModel("gemini-1.5-flash")
                response = model.generate_content(full_prompt)
                ai_sql = response.text.strip()
                ai_sql = ai_sql.replace("```sql", "").replace("```", "").strip()

                st.session_state["ai_generated_sql"] = ai_sql

                st.write("### Generated SQL Query")
                st.code(ai_sql, language="sql")

            except Exception as e:
                st.error(f"AI Error: {e}")

    if "ai_generated_sql" in st.session_state:
        if st.button("Execute AI Query"):
            try:
                result = pd.read_sql(
                    st.session_state["ai_generated_sql"],
                    conn
                )
                st.session_state.query_history.append(
                    st.session_state["ai_generated_sql"]
                )
                st.write("### Result")
                st.dataframe(result, use_container_width=True)
                del st.session_state["ai_generated_sql"]
            except Exception as e:
                st.error(f"SQL Error: {e}")